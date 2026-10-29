from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from PySide6.QtCore import Signal
from PySide6.QtWidgets import (
    QComboBox, QFileDialog, QHBoxLayout, QLabel, QMessageBox, QPushButton,
    QTabWidget, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

from .settings import app_data_dir


SOURCE_NAMES = ("BAL BROUGHT FORWARD", "VL", "SL", "FL")


def read_workbook(path: str) -> dict[str, list[list[object]]]:
    """Read workbook values without changing the source file."""
    suffix = Path(path).suffix.casefold()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(path)
        result = {}
        for sheet in workbook.sheets():
            rows = []
            for row_index in range(sheet.nrows):
                row = []
                for cell in sheet.row(row_index):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        row.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
                    else:
                        row.append(cell.value)
                rows.append(row)
            result[sheet.name] = rows
        return result
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
                    for sheet in workbook.worksheets}
        finally:
            workbook.close()
    raise ValueError("Choose an .xls or .xlsx workbook.")


def display_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time().isoformat() == "00:00:00" else str(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


class Ver2SourceTab(QWidget):
    source_changed = Signal()

    def __init__(self, title: str) -> None:
        super().__init__()
        self.path = ""
        self.sheets: dict[str, list[list[object]]] = {}
        layout = QVBoxLayout(self)
        controls = QHBoxLayout()
        button = QPushButton(f"Choose {title} Excel file")
        button.clicked.connect(self.choose_file)
        controls.addWidget(button)
        self.file_label = QLabel("No file selected")
        self.file_label.setWordWrap(True)
        controls.addWidget(self.file_label, 1)
        controls.addWidget(QLabel("Sheet:"))
        self.sheet_selector = QComboBox()
        self.sheet_selector.currentTextChanged.connect(self._sheet_changed)
        controls.addWidget(self.sheet_selector)
        refresh = QPushButton("Reload file")
        refresh.clicked.connect(self.reload)
        controls.addWidget(refresh)
        layout.addLayout(controls)
        self.status = QLabel("Choose a workbook to preview its contents.")
        self.status.setWordWrap(True)
        layout.addWidget(self.status)
        self.table = QTableWidget()
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setStyleSheet("QTableWidget { font-size: 14pt; } QHeaderView::section { font-size: 14pt; padding: 5px; }")
        layout.addWidget(self.table, 1)

    def choose_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Select Excel file", "", "Excel files (*.xls *.xlsx)")
        if path:
            self.load_file(path)

    def reload(self) -> None:
        if self.path:
            self.load_file(self.path, self.sheet_selector.currentText())

    def load_file(self, path: str, selected_sheet: str = "", *, notify: bool = True) -> None:
        try:
            sheets = read_workbook(path)
            if not sheets:
                raise ValueError("The workbook has no sheets.")
        except Exception as error:
            self.status.setText(f"Could not load {Path(path).name}: {error}")
            if notify:
                QMessageBox.warning(self, "Excel file", self.status.text())
            return
        self.path, self.sheets = path, sheets
        self.file_label.setText(path)
        self.sheet_selector.blockSignals(True)
        self.sheet_selector.clear()
        self.sheet_selector.addItems(list(sheets))
        if selected_sheet in sheets:
            self.sheet_selector.setCurrentText(selected_sheet)
        self.sheet_selector.blockSignals(False)
        self._render_sheet()
        if notify:
            self.source_changed.emit()

    def _sheet_changed(self, _name: str) -> None:
        self._render_sheet()
        self.source_changed.emit()

    def _render_sheet(self) -> None:
        rows = self.sheets.get(self.sheet_selector.currentText(), [])
        width = max((len(row) for row in rows), default=0)
        self.table.clear()
        self.table.setRowCount(len(rows))
        self.table.setColumnCount(width)
        self.table.setHorizontalHeaderLabels([str(i + 1) for i in range(width)])
        for row_index, row in enumerate(rows):
            for column, value in enumerate(row):
                text = display_value(value)
                if text:
                    self.table.setItem(row_index, column, QTableWidgetItem(text))
        self.status.setText(f"Preview: {len(rows)} rows · {width} columns. Source file is unchanged.")


class Ver2Page(QWidget):
    back_requested = Signal()

    def __init__(self) -> None:
        super().__init__()
        layout = QVBoxLayout(self)
        heading = QHBoxLayout()
        title = QLabel("VER2 · Excel Leave Data")
        title.setStyleSheet("font-size: 22pt; font-weight: 800;")
        heading.addWidget(title)
        heading.addStretch(1)
        back = QPushButton("← VER1")
        back.clicked.connect(self.back_requested.emit)
        heading.addWidget(back)
        layout.addLayout(heading)
        note = QLabel("Select a workbook for each source. This preview does not import records or calculate balances yet.")
        note.setWordWrap(True)
        layout.addWidget(note)
        self.tabs = QTabWidget()
        self.sources = {}
        for name in SOURCE_NAMES:
            tab = Ver2SourceTab(name)
            tab.source_changed.connect(self.save_sources)
            self.sources[name] = tab
            self.tabs.addTab(tab, name)
        layout.addWidget(self.tabs, 1)
        self._sources_restored = False

    def restore_sources_once(self) -> None:
        if not self._sources_restored:
            self._sources_restored = True
            self._restore_sources()

    def _restore_sources(self) -> None:
        config = app_data_dir() / "ver2_sources.json"
        try:
            data = json.loads(config.read_text(encoding="utf-8")) if config.exists() else {}
        except (OSError, ValueError):
            data = {}
        if not isinstance(data, dict):
            return
        for name, tab in self.sources.items():
            source = data.get(name, {})
            if isinstance(source, dict) and source.get("path"):
                tab.load_file(str(source["path"]), str(source.get("sheet", "")), notify=False)

    def save_sources(self) -> None:
        config = app_data_dir() / "ver2_sources.json"
        data = {name: {"path": tab.path, "sheet": tab.sheet_selector.currentText()}
                for name, tab in self.sources.items() if tab.path}
        try:
            config.write_text(json.dumps(data, indent=2), encoding="utf-8")
        except OSError as error:
            QMessageBox.warning(self, "VER2 sources", f"Could not save source selection: {error}")
